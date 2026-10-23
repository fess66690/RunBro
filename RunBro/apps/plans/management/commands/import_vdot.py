"""
Импорт VDOT-справочника из Excel.
Читает: data/raw/VDOT.xlsx (прогнозы) + data/raw/TrainingTemp.xlsx (темпы).
"""
from datetime import time, timedelta
from pathlib import Path

import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from apps.plans.models import VdotReference


def pace_to_seconds(value) -> int:
    """
    Темпы (Easy, Marathon, Threshold, Interval).
    Excel сохраняет '5:46' как '05:46:00' (часы=5, минуты=46).
    По смыслу это 5 мин 46 сек = 346 сек.
    """
    if value is None:
        return 0
    if isinstance(value, time):
        return value.hour * 60 + value.minute  # 5*60 + 46 = 346
    if isinstance(value, timedelta):
        return int(value.total_seconds())
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str):
        parts = value.strip().split(':')
        try:
            nums = [int(p) for p in parts]
        except ValueError:
            return 0
        if len(nums) == 2:
            return nums[0] * 60 + nums[1]
        if len(nums) == 3:
            return nums[0] * 60 + nums[1]  # часы трактуем как минуты
    return 0


def time_to_seconds(value, hours_as_hours=False) -> int:
    """
    Прогнозы на гонки.
    hours_as_hours=True — для длинных (полумарафон, марафон): 2:21:04 → 8464
    hours_as_hours=False — для коротких (1000м…10к): 3:48:00 → 228 сек
    """
    if value is None:
        return 0
    if isinstance(value, time):
        if hours_as_hours:
            return value.hour * 3600 + value.minute * 60 + value.second
        return value.hour * 60 + value.minute
    if isinstance(value, timedelta):
        return int(value.total_seconds())
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str):
        parts = value.strip().split(':')
        try:
            nums = [int(p) for p in parts]
        except ValueError:
            return 0
        if len(nums) == 2:
            return nums[0] * 60 + nums[1]
        if len(nums) == 3:
            if hours_as_hours:
                return nums[0] * 3600 + nums[1] * 60 + nums[2]
            return nums[0] * 60 + nums[1]
    return 0


class Command(BaseCommand):
    help = 'Импорт VDOT-справочника из Excel (VDOT.xlsx + TrainingTemp.xlsx)'

    def add_arguments(self, parser):
        parser.add_argument('--vdot-file', type=str, default='data/raw/VDOT.xlsx')
        parser.add_argument('--temp-file', type=str, default='data/raw/TrainingTemp.xlsx')
        parser.add_argument('--clear', action='store_true', help='Удалить все записи перед импортом')

    def handle(self, *args, **options):
        vdot_path = Path(options['vdot_file'])
        temp_path = Path(options['temp_file'])
        if not vdot_path.is_absolute():
            vdot_path = Path(settings.BASE_DIR) / vdot_path
        if not temp_path.is_absolute():
            temp_path = Path(settings.BASE_DIR) / temp_path

        if not vdot_path.exists():
            raise CommandError(f'Файл не найден: {vdot_path}')
        if not temp_path.exists():
            raise CommandError(f'Файл не найден: {temp_path}')

        self.stdout.write(f'[dir] VDOT:         {vdot_path}')
        self.stdout.write(f'[dir] TrainingTemp: {temp_path}\n')

        # --- 1. VDOT.xlsx (прогнозы) ---
        vdot_wb = openpyxl.load_workbook(vdot_path, data_only=True)
        vdot_ws = vdot_wb.active
        vdot_rows = list(vdot_ws.iter_rows(values_only=True))
        self.stdout.write(f'   VDOT.xlsx колонки: {vdot_rows[0]}')

        vdot_data = {}
        for row in vdot_rows[1:]:
            if row[0] is None:
                continue
            try:
                vdot_value = float(row[0])
            except (TypeError, ValueError):
                continue
            vdot_data[vdot_value] = {
                # Короткие дистанции — MM:SS
                'predicted_1000m_sec': time_to_seconds(row[1], hours_as_hours=False),
                'predicted_1500m_sec': time_to_seconds(row[2], hours_as_hours=False),
                'predicted_mile_sec': time_to_seconds(row[3], hours_as_hours=False),
                'predicted_3000m_sec': time_to_seconds(row[4], hours_as_hours=False),
                'predicted_2mile_sec': time_to_seconds(row[5], hours_as_hours=False),
                'predicted_5k_sec': time_to_seconds(row[6], hours_as_hours=False),
                'predicted_10k_sec': time_to_seconds(row[7], hours_as_hours=False),
                # Длинные — HH:MM:SS
                'predicted_half_sec': time_to_seconds(row[8], hours_as_hours=True),
                'predicted_marathon_sec': time_to_seconds(row[9], hours_as_hours=True),
            }
        self.stdout.write(f'   [ok] Прогнозов: {len(vdot_data)}\n')

        # --- 2. TrainingTemp.xlsx (темпы) ---
        temp_wb = openpyxl.load_workbook(temp_path, data_only=True)
        temp_ws = temp_wb.active
        temp_rows = list(temp_ws.iter_rows(values_only=True))
        self.stdout.write(f'   TrainingTemp.xlsx колонки: {temp_rows[0]}')

        temp_data = {}
        for row in temp_rows[1:]:
            if row[0] is None:
                continue
            try:
                vdot_value = float(row[0])
            except (TypeError, ValueError):
                continue
            temp_data[vdot_value] = {
                'easy_pace_sec': pace_to_seconds(row[1]),
                'marathon_pace_sec': pace_to_seconds(row[2]),
                'threshold_pace_sec': pace_to_seconds(row[3]),
                'interval_pace_sec': pace_to_seconds(row[4]),
                'repetition_pace_200m_sec': int(row[5]) if row[5] else 0,
                'fast_reps_pace_200m_sec': int(row[6]) if row[6] else 0,
            }
        self.stdout.write(f'   [ok] Темпов: {len(temp_data)}\n')

        # --- 3. Сохраняем ---
        all_vdots = sorted(set(vdot_data.keys()) | set(temp_data.keys()))
        self.stdout.write(f'[merge] Объединяем: {len(all_vdots)} уникальных VDOT')

        with transaction.atomic():
            if options['clear']:
                deleted, _ = VdotReference.objects.all().delete()
                self.stdout.write(f'[deleted] Удалено старых: {deleted}')

            created = updated = 0
            for vdot_value in all_vdots:
                data = {}
                data.update(vdot_data.get(vdot_value, {}))
                data.update(temp_data.get(vdot_value, {}))

                _, is_created = VdotReference.objects.update_or_create(
                    vdot_value=vdot_value,
                    defaults=data,
                )
                if is_created:
                    created += 1
                else:
                    updated += 1

        self.stdout.write(self.style.SUCCESS(
            f'\n[done] Готово! Создано: {created}, обновлено: {updated}'
        ))
        self.stdout.write(f'   Всего в БД: {VdotReference.objects.count()}')